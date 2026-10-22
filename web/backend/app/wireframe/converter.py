from __future__ import annotations

import html
import math
import re
from dataclasses import dataclass, field
from typing import BinaryIO

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_COLOR_TYPE, MSO_FILL
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE as ST
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

EMU_PER_PX = 9525
CALLOUT_RE = re.compile(r"^#(\d+)$")
PAGE_ID_RE = re.compile(r"^[A-Z]+(-[A-Z0-9]+)+$")
HEADER_MAX_TOP_PX = 30
SCREEN_WIDTH_PX = 962
NEUTRAL_FILL = "#eef1f2"
NEUTRAL_LINE = "#9aa8ab"


@dataclass
class Annotation:
    number: int
    text: str


@dataclass
class SlideDoc:
    index: int
    project: str
    title: str
    page_id: str
    width: int
    height: int
    body: str
    annotations: list[Annotation] = field(default_factory=list)
    items: list[tuple] = field(default_factory=list)


def _px(emu: float) -> float:
    return emu / EMU_PER_PX


def _rgb(color) -> str | None:
    try:
        if color.type == MSO_COLOR_TYPE.RGB:
            return f"#{color.rgb}"
    except (AttributeError, TypeError):
        pass
    return None


def _soften(hex_color: str) -> str:
    r, g, b = (int(hex_color[i:i + 2], 16) for i in (1, 3, 5))
    luma = 0.299 * r + 0.587 * g + 0.114 * b
    gray = round(255 - (255 - luma) * 0.45)
    return f"#{gray:02x}{gray:02x}{gray:02x}"


def _fill_of(shape) -> str | None:
    try:
        fill = shape.fill
        kind = fill.type
    except (AttributeError, TypeError, NotImplementedError):
        return None
    if kind == MSO_FILL.SOLID:
        rgb = _rgb(fill.fore_color)
        return _soften(rgb) if rgb else NEUTRAL_FILL
    if kind == MSO_FILL.PICTURE:
        return "image"
    if kind in (MSO_FILL.GRADIENT, MSO_FILL.PATTERNED):
        return NEUTRAL_FILL
    if kind is None and shape.shape_type == ST.AUTO_SHAPE:
        return NEUTRAL_FILL
    return None


def _border_of(shape) -> tuple[str, float] | None:
    try:
        line = shape.line
        kind = line.fill.type
        width = _px(line.width) if line.width else 1.0
    except (AttributeError, TypeError, NotImplementedError):
        return None
    if kind == MSO_FILL.SOLID:
        rgb = _rgb(line.fill.fore_color)
        return (_soften(rgb) if rgb else NEUTRAL_LINE), max(1.0, min(width, 3.0))
    if kind is None and shape.shape_type == ST.AUTO_SHAPE:
        return NEUTRAL_LINE, 1.0
    return None


def _radius_of(shape) -> str:
    try:
        kind = shape.auto_shape_type
    except (AttributeError, ValueError, NotImplementedError):
        return "0"
    if kind == MSO_SHAPE.OVAL:
        return "50%"
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        return "8px"
    return "0"


_ALIGN = {PP_ALIGN.CENTER: "center", PP_ALIGN.RIGHT: "right", PP_ALIGN.JUSTIFY: "justify"}
_ANCHOR = {MSO_ANCHOR.MIDDLE: "center", MSO_ANCHOR.BOTTOM: "flex-end"}


def _text_html(shape) -> tuple[str, str]:
    if not shape.has_text_frame:
        return "", ""
    frame = shape.text_frame
    paragraphs = []
    for paragraph in frame.paragraphs:
        text = paragraph.text.strip()
        size = next((r.font.size.pt for r in paragraph.runs if r.font.size), None)
        if size is None and paragraph.font.size:
            size = paragraph.font.size.pt
        font_px = (size or 11) * 4 / 3
        style = [f"font-size:{font_px:.1f}px"]
        if any(r.font.bold for r in paragraph.runs):
            style.append("font-weight:700")
        align = _ALIGN.get(paragraph.alignment)
        if align:
            style.append(f"text-align:{align}")
        spacing = paragraph.line_spacing
        if isinstance(spacing, float):
            style.append(f"line-height:{spacing * 1.2:.2f}")
        elif spacing is not None:
            style.append(f"line-height:{_px(spacing):.1f}px")
        if paragraph.space_before:
            style.append(f"margin-top:{_px(paragraph.space_before):.1f}px")
        if paragraph.space_after:
            style.append(f"margin-bottom:{_px(paragraph.space_after):.1f}px")
        body = html.escape(text).replace("\v", "<br>") if text else "&nbsp;"
        paragraphs.append(f'<div style="{";".join(style)}">{body}</div>')
    if not any(p.strip() for p in frame.text.split("\n")):
        return "", ""
    anchor = _ANCHOR.get(frame.vertical_anchor, "flex-start")
    return "".join(paragraphs), anchor


def _plain_text(shape) -> str:
    return shape.text_frame.text.strip() if shape.has_text_frame else ""


@dataclass
class _Transform:
    ox: float = 0.0
    oy: float = 0.0
    sx: float = 1.0
    sy: float = 1.0

    def point(self, x: float, y: float) -> tuple[float, float]:
        return self.ox + self.sx * x, self.oy + self.sy * y

    def child(self, group) -> "_Transform":
        xfrm = group._element.grpSpPr.xfrm
        if xfrm is None or not xfrm.chExt.cx or not xfrm.chExt.cy:
            return self
        sx2 = self.sx * xfrm.ext.cx / xfrm.chExt.cx
        sy2 = self.sy * xfrm.ext.cy / xfrm.chExt.cy
        return _Transform(
            self.ox + self.sx * xfrm.off.x - sx2 * xfrm.chOff.x,
            self.oy + self.sy * xfrm.off.y - sy2 * xfrm.chOff.y,
            sx2,
            sy2,
        )


@dataclass
class _Canvas:
    origin_y: float
    parts: list[str] = field(default_factory=list)
    lines: list[str] = field(default_factory=list)
    items: list[tuple] = field(default_factory=list)

    def abs_style(self, x, y, w, h) -> str:
        return f"left:{x:.1f}px;top:{y - self.origin_y:.1f}px;width:{w:.1f}px;height:{h:.1f}px"


def _emit_table(shape, canvas: _Canvas, tf: _Transform) -> None:
    x, y = tf.point(shape.left, shape.top)
    table = shape.table
    widths = [_px(c.width) * tf.sx for c in table.columns]
    rows = []
    for row in table.rows:
        cells = "".join(
            f"<td>{html.escape(cell.text.strip()).replace(chr(10), '<br>')}</td>"
            for cell in row.cells
        )
        rows.append(f'<tr style="height:{_px(row.height) * tf.sy:.1f}px">{cells}</tr>')
    cols = "".join(f'<col style="width:{w:.1f}px">' for w in widths)
    canvas.items.append((
        "table", _px(x), _px(y) - canvas.origin_y, sum(widths), sum(_px(r.height) * tf.sy for r in table.rows),
        " / ".join(" | ".join(c.text.strip().replace("\n", " ") for c in r.cells) for r in table.rows),
    ))
    canvas.parts.append(
        f'<div class="wf-table" style="left:{_px(x):.1f}px;top:{_px(y) - canvas.origin_y:.1f}px">'
        f"<table><colgroup>{cols}</colgroup>{''.join(rows)}</table></div>"
    )


def _connector_points(shape, tf: _Transform) -> list[tuple[float, float]]:
    xfrm = shape._element.spPr.xfrm
    x, y = tf.point(shape.left, shape.top)
    w, h = shape.width * tf.sx, shape.height * tf.sy
    begin_x = w if xfrm.get("flipH") == "1" else 0.0
    begin_y = h if xfrm.get("flipV") == "1" else 0.0
    end_x, end_y = w - begin_x, h - begin_y
    preset = (shape._element.xpath(".//a:prstGeom/@prst") or [""])[0]
    if preset == "bentConnector2":
        points = [(begin_x, begin_y), (end_x, begin_y), (end_x, end_y)]
    elif preset.startswith("bentConnector"):
        mid_x = (begin_x + end_x) / 2
        points = [(begin_x, begin_y), (mid_x, begin_y), (mid_x, end_y), (end_x, end_y)]
    else:
        points = [(begin_x, begin_y), (end_x, end_y)]
    angle = math.radians(int(xfrm.get("rot", "0")) / 60000)
    cos, sin = math.cos(angle), math.sin(angle)
    cx, cy = w / 2, h / 2
    return [
        (x + cx + (px - cx) * cos - (py - cy) * sin, y + cy + (px - cx) * sin + (py - cy) * cos)
        for px, py in points
    ]


def _emit_line(shape, canvas: _Canvas, tf: _Transform) -> None:
    points = " ".join(
        f"{_px(px):.1f},{_px(py) - canvas.origin_y:.1f}" for px, py in _connector_points(shape, tf)
    )
    dashed = ' stroke-dasharray="5 4"' if shape._element.xpath(".//a:prstDash[@val!='solid']") else ""
    start = ' marker-start="url(#wf-arrow)"' if shape._element.xpath(".//a:headEnd[@type!='none']") else ""
    end = ' marker-end="url(#wf-arrow)"' if shape._element.xpath(".//a:tailEnd[@type!='none']") else ""
    canvas.lines.append(
        f'<polyline points="{points}" fill="none" stroke="#8a9a9d" stroke-width="1.5"{dashed}{start}{end}/>'
    )


def _emit_shape(shape, canvas: _Canvas, tf: _Transform) -> None:
    x, y = tf.point(shape.left, shape.top)
    w, h = shape.width * tf.sx, shape.height * tf.sy
    px, py, pw, ph = _px(x), _px(y), _px(w), _px(h)
    plain = _plain_text(shape)
    callout = CALLOUT_RE.match(plain)
    if callout and pw <= 40 and ph <= 40:
        canvas.parts.append(
            f'<div class="wf-badge" data-n="{callout.group(1)}" style="{canvas.abs_style(px, py, pw, ph)}">'
            f"{callout.group(1)}</div>"
        )
        return
    fill = _fill_of(shape)
    border = _border_of(shape)
    is_picture = shape.shape_type == ST.PICTURE
    text, anchor = _text_html(shape)
    if not (fill or border or is_picture or text):
        return
    kind = "image" if (is_picture or fill == "image") else ("box" if (fill or border) else "text")
    canvas.items.append((kind, px, py - canvas.origin_y, pw, ph, plain.replace("\v", " ").replace("\n", " / ")))
    style = [canvas.abs_style(px, py, pw, ph)]
    classes = ["wf-box"]
    if not (fill or border or is_picture):
        classes.append("wf-text")
    if is_picture or fill == "image":
        classes.append("wf-image")
    elif fill:
        style.append(f"background:{fill}")
    if border:
        style.append(f"border:{border[1]:.1f}px solid {border[0]}")
    radius = _radius_of(shape)
    if radius != "0":
        style.append(f"border-radius:{radius}")
    if shape.rotation:
        style.append(f"transform:rotate({shape.rotation:.1f}deg)")
    if anchor:
        style.append(f"justify-content:{anchor}")
    canvas.parts.append(f'<div class="{" ".join(classes)}" style="{";".join(style)}">{text}</div>')


def _walk(shapes, canvas: _Canvas, tf: _Transform, skip) -> None:
    for shape in shapes:
        if skip(shape, tf):
            continue
        kind = shape.shape_type
        if kind == ST.GROUP:
            _walk(shape.shapes, canvas, tf.child(shape), skip)
        elif kind == ST.TABLE:
            _emit_table(shape, canvas, tf)
        elif kind == ST.LINE or shape._element.tag.endswith("}cxnSp"):
            _emit_line(shape, canvas, tf)
        else:
            _emit_shape(shape, canvas, tf)


def _annotations(table) -> list[Annotation]:
    found = []
    for row in table.rows:
        cells = [c.text.strip() for c in row.cells]
        match = CALLOUT_RE.match(cells[0]) if cells else None
        text = "\n".join(c for c in cells[1:] if c)
        if match and text:
            found.append(Annotation(int(match.group(1)), text))
    return found


def _has_callouts(table) -> bool:
    return any(CALLOUT_RE.match(row.cells[0].text.strip()) for row in table.rows)


def _is_header(shape) -> bool:
    return shape.shape_type == ST.PLACEHOLDER and _px(shape.top) < HEADER_MAX_TOP_PX


def _is_page_number(shape) -> bool:
    return shape.shape_type == ST.PLACEHOLDER and _px(shape.top) > 600 and _plain_text(shape).isdigit()


def _split_header(shapes) -> tuple[str, str, str]:
    texts = [t for t in (_plain_text(s) for s in shapes if _is_header(s)) if t and t != "-"]
    project = next((t for t in texts if t.startswith("[")), "")
    page_id = next((t for t in texts if PAGE_ID_RE.match(t)), "")
    title = next((t for t in texts if t not in (project, page_id)), "")
    return project, title, page_id


def convert(source: str | BinaryIO) -> list[SlideDoc]:
    presentation = Presentation(source)
    slide_w = _px(presentation.slide_width)
    slide_h = _px(presentation.slide_height)
    docs = []
    for index, slide in enumerate(presentation.slides, 1):
        panel = next(
            (s for s in slide.shapes if s.shape_type == ST.TABLE and _has_callouts(s.table)),
            None,
        )
        annotations = _annotations(panel.table) if panel is not None else []
        project, title, page_id = _split_header(slide.shapes)
        has_header = bool(project or page_id)
        origin_y = 26.0 if has_header else 0.0
        canvas = _Canvas(origin_y=origin_y)

        def skip(shape, tf, panel=panel, has_header=has_header):
            if (panel is not None and shape._element is panel._element) or _is_page_number(shape):
                return True
            return has_header and _is_header(shape)

        _walk(slide.shapes, canvas, _Transform(), skip)
        width = SCREEN_WIDTH_PX if panel is not None else slide_w
        height = slide_h - origin_y
        svg = ""
        if canvas.lines:
            svg = (
                f'<svg class="wf-lines" width="{width:.0f}" height="{height:.0f}">'
                '<defs><marker id="wf-arrow" markerWidth="8" markerHeight="8" refX="7" refY="4" orient="auto-start-reverse">'
                '<path d="M0,0 L8,4 L0,8 z" fill="#8a9a9d"/></marker></defs>'
                f'{"".join(canvas.lines)}</svg>'
            )
        docs.append(SlideDoc(index, project, title, page_id, int(width), int(height), "".join(canvas.parts) + svg, annotations, canvas.items))
    return docs
