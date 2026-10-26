import io
import sys
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import pytest
from pptx import Presentation
from pptx.util import Emu

import app.deps as deps_module
import app.wireframe.routes as wireframe_routes


def _screen_deck() -> bytes:
    prs = Presentation()
    prs.slide_width = Emu(12192000)
    prs.slide_height = Emu(6858000)
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "DN-SEA-P-0"
    slide.shapes.title.top = Emu(100000)
    slide.shapes.title.height = Emu(200000)
    panel = slide.shapes.add_table(2, 2, Emu(9163900), Emu(472800), Emu(3024000), Emu(600000))
    panel.table.cell(0, 0).text = "#1"
    panel.table.cell(0, 1).text = "로고 - 클릭 시 홈으로 이동"
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


@pytest.fixture(autouse=True)
def _isolate_job_side_effects(db_factory, tmp_path, monkeypatch):
    # 생성 작업 스레드가 실제 로컬 DB와 data/ 폴더를 건드리지 않게 테스트 DB·임시 폴더로 돌린다.
    monkeypatch.setattr(wireframe_routes, "SessionLocal", db_factory)
    monkeypatch.setattr(wireframe_routes, "_WIRE_OUTPUT_DIR", tmp_path)


@pytest.fixture(autouse=True)
def _reset_state():
    wireframe_routes._jobs.clear()
    wireframe_routes._themes.clear()


@pytest.fixture()
def admin(client, make_user, monkeypatch):
    monkeypatch.setattr(deps_module, "ADMIN_EMAIL", "admin@team.com")
    make_user(email="admin@team.com", password="correcthorse", name="관리자")
    client.post("/auth/login", json={"email": "admin@team.com", "password": "correcthorse"})


def _upload(client, path, name="screen.pptx", content=None, **data):
    return client.post(
        path,
        files={"file": (name, content if content is not None else _screen_deck())},
        data=data,
    )


def _finish(client, job_id):
    for _ in range(400):
        body = client.get(f"/wireframe/jobs/{job_id}").json()
        if body["status"] in ("done", "error"):
            return body
        time.sleep(0.05)
    raise AssertionError("job did not finish")


def test_wireframe_requires_login(client):
    assert _upload(client, "/wireframe/analyze").status_code == 401


def test_wireframe_rejects_non_admin(client, make_user, monkeypatch):
    monkeypatch.setattr(deps_module, "ADMIN_EMAIL", "admin@team.com")
    make_user(email="other@team.com", password="hunter2", name="다른사람")
    client.post("/auth/login", json={"email": "other@team.com", "password": "hunter2"})
    assert _upload(client, "/wireframe/analyze").status_code == 403


def test_analyze_lists_screen_pages(client, admin):
    res = _upload(client, "/wireframe/analyze")
    assert res.status_code == 200
    assert res.json()["pages"] == [{"pageId": "DN-SEA-P-0", "title": "", "slides": [1], "elementCount": 0}]


def test_analyze_rejects_non_pptx_and_unreadable_files(client, admin):
    assert _upload(client, "/wireframe/analyze", name="a.xlsx").status_code == 400
    assert _upload(client, "/wireframe/analyze", content=b"not a pptx").status_code == 400


def test_generate_rejects_unknown_page(client, admin):
    assert _upload(client, "/wireframe/generate", page_id="NOPE").status_code == 404


def test_job_generates_theme_once_per_deck(client, admin, monkeypatch):
    theme_calls = []

    def fake_theme(context, on_text=None):
        theme_calls.append(context)
        return ":root{--bg:#fff}", {"input": 1000, "output": 1000}

    monkeypatch.setattr(wireframe_routes, "generate_theme", fake_theme)
    monkeypatch.setattr(
        wireframe_routes,
        "generate_page_hifi",
        lambda page, context, theme_css, related, instruction="", on_text=None: (
            f"<!doctype html><html>{theme_css}</html>",
            {"input": 2000, "output": 3000, "cacheRead": 0, "cacheWrite": 0},
        ),
    )
    first = _finish(client, _upload(client, "/wireframe/generate", page_id="DN-SEA-P-0").json()["jobId"])
    second = _finish(client, _upload(client, "/wireframe/generate", page_id="DN-SEA-P-0").json()["jobId"])

    assert first["status"] == second["status"] == "done"
    assert len(theme_calls) == 1
    assert ":root{--bg:#fff}" in first["result"]["html"]
    # 첫 호출만 테마 비용(입력 1000, 출력 1000)을 함께 청구한다
    assert first["result"]["costUsd"] > second["result"]["costUsd"]
    assert second["result"]["costUsd"] == pytest.approx((2000 * 2 + 3000 * 10) / 1e6)


def test_job_reports_ai_failure(client, admin, monkeypatch):
    def boom(context, on_text=None):
        raise RuntimeError("credit balance is too low")

    monkeypatch.setattr(wireframe_routes, "generate_theme", boom)
    res = _upload(client, "/wireframe/generate", page_id="DN-SEA-P-0")
    body = _finish(client, res.json()["jobId"])
    assert body["status"] == "error"
    assert "credit balance" in body["error"]


def test_unknown_job_is_404(client, admin):
    assert client.get("/wireframe/jobs/nope").status_code == 404
