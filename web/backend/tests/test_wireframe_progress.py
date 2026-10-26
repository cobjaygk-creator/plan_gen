import io
import sys
import threading
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import pytest
from pptx import Presentation
from pptx.util import Emu

import app.deps as deps_module
import app.wireframe.routes as wireframe_routes
from app.wireframe.routes import _progress


def _deck() -> bytes:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "DN-SEA-P-0"
    slide.shapes.title.top, slide.shapes.title.height = Emu(100000), Emu(200000)
    panel = slide.shapes.add_table(2, 2, Emu(9163900), Emu(472800), Emu(3024000), Emu(600000))
    panel.table.cell(0, 0).text = "#1"
    panel.table.cell(0, 1).text = "로고 - 클릭 시 홈으로 이동"
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def _job(**fields):
    now = time.time()
    return {"status": "running", "stage": "queued", "stageStartedAt": now, "startedAt": now, "chars": 0, "progress": 0, **fields}


def test_progress_grows_with_streamed_text_inside_a_stage():
    low = _progress(_job(stage="page", themeGenerated=True, chars=0))
    high = _progress(_job(stage="page", themeGenerated=True, chars=11_000))
    assert 30 <= low < high < 96


def test_progress_never_goes_backwards_and_caps_before_done():
    job = _job(stage="page", themeGenerated=True, chars=500_000)
    first = _progress(job)
    job["stage"], job["chars"] = "theme", 0
    assert _progress(job) >= first
    assert first <= 96 - 1


def test_page_stage_starts_earlier_when_the_theme_is_reused():
    reused = _progress(_job(stage="page", themeGenerated=False))
    fresh = _progress(_job(stage="page", themeGenerated=True))
    assert reused < fresh and reused >= 8


def test_done_job_reports_100():
    assert _progress(_job(status="done", stage="saving")) == 100


@pytest.fixture(autouse=True)
def _isolate(db_factory, tmp_path, monkeypatch):
    monkeypatch.setattr(wireframe_routes, "SessionLocal", db_factory)
    monkeypatch.setattr(wireframe_routes, "_WIRE_OUTPUT_DIR", tmp_path)
    wireframe_routes._jobs.clear()
    wireframe_routes._themes.clear()


@pytest.fixture()
def admin(client, make_user, monkeypatch):
    monkeypatch.setattr(deps_module, "ADMIN_EMAIL", "admin@team.com")
    make_user(email="admin@team.com", password="correcthorse", name="관리자")
    client.post("/auth/login", json={"email": "admin@team.com", "password": "correcthorse"})


def test_api_reports_stage_and_rising_progress_while_the_theme_is_generated(client, admin, monkeypatch):
    release = threading.Event()

    def slow_theme(context, on_text=None):
        on_text(5_000)
        release.wait(10)
        return ":root{}", {"input": 1, "output": 1}

    monkeypatch.setattr(wireframe_routes, "generate_theme", slow_theme)
    monkeypatch.setattr(
        wireframe_routes, "generate_page_hifi",
        lambda page, context, theme_css, related, instruction="", on_text=None: (
            "<!doctype html><html></html>", {"input": 1, "output": 1, "cacheRead": 0, "cacheWrite": 0},
        ),
    )
    job_id = client.post("/wireframe/generate", files={"file": ("sb.pptx", _deck())}, data={"page_id": "DN-SEA-P-0"}).json()["jobId"]

    seen = None
    for _ in range(200):
        body = client.get(f"/wireframe/jobs/{job_id}").json()
        if body["stage"] == "theme" and body["progress"] > 5:
            seen = body
            break
        time.sleep(0.05)
    release.set()
    assert seen is not None and seen["status"] == "running"
    assert 5 < seen["progress"] < 30 and seen["elapsedSec"] >= 0

    for _ in range(400):
        final = client.get(f"/wireframe/jobs/{job_id}").json()
        if final["status"] == "done":
            break
        time.sleep(0.05)
    assert final["status"] == "done" and final["progress"] == 100 and final["result"]["html"].startswith("<!doctype html>")
