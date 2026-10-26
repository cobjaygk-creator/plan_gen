import io
import sys
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import pytest
from pptx import Presentation
from pptx.util import Emu

import app.deps as deps_module
import app.wireframe.routes as wireframe_routes
from app.wireframe.converter import convert
from app.wireframe.hifi import deck_context, find_settings_slide

GUIDE_TEXT = [
    "SB 첫 페이지 가이드",
    "01 화면 범위와 목적",
    "화면 범위: 슬라이드 2~3",
    "03 레이아웃과 정보 구조",
    "레이아웃: 좌측 사이드바형",
    "05 시각 기준과 제약",
    "스타일: 모던 SaaS형 / 해상도 1440×900",
]


def _add_text_slide(prs, lines):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for i, line in enumerate(lines):
        box = slide.shapes.add_textbox(Emu(300000), Emu(300000 + i * 400000), Emu(8000000), Emu(350000))
        box.text_frame.text = line
    return slide


def _add_screen_slide(prs, page_id="DN-SEA-P-0"):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = page_id
    slide.shapes.title.top = Emu(100000)
    slide.shapes.title.height = Emu(200000)
    panel = slide.shapes.add_table(2, 2, Emu(9163900), Emu(472800), Emu(3024000), Emu(600000))
    panel.table.cell(0, 0).text = "#1"
    panel.table.cell(0, 1).text = "로고 - 클릭 시 홈으로 이동"
    return slide


def _deck(*builders) -> bytes:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    for build in builders:
        build(prs)
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def _docs(*builders):
    return convert(io.BytesIO(_deck(*builders)))


def test_guide_headings_identify_the_settings_slide():
    docs = _docs(lambda p: _add_text_slide(p, ["표지"]), lambda p: _add_text_slide(p, GUIDE_TEXT), _add_screen_slide)
    assert find_settings_slide(docs).index == 2


def test_keywords_need_the_word_wireframe():
    keywords_only = ["레이아웃: 사이드바", "스타일: 모던", "필수: 검색"]
    assert find_settings_slide(_docs(lambda p: _add_text_slide(p, keywords_only))) is None
    with_word = ["와이어프레임 설정", *keywords_only]
    assert find_settings_slide(_docs(lambda p: _add_text_slide(p, with_word))) is not None


def test_screen_slides_and_late_slides_are_not_settings():
    docs = _docs(_add_screen_slide)
    assert find_settings_slide(docs) is None
    filler = [lambda p: _add_text_slide(p, ["본문"]) for _ in range(6)]
    late = _docs(*filler, lambda p: _add_text_slide(p, GUIDE_TEXT))
    assert find_settings_slide(late) is None


def test_deck_context_puts_settings_first_and_does_not_repeat_it():
    docs = _docs(lambda p: _add_text_slide(p, ["업무 개요", "작업 사이즈 2560"]), lambda p: _add_text_slide(p, GUIDE_TEXT))
    context = deck_context(docs)
    assert context.startswith("[와이어프레임 설정 장표 — 최우선 반영]")
    assert "스타일: 모던 SaaS형" in context
    assert context.count("스타일: 모던 SaaS형") == 1
    assert "업무 개요" in context


@pytest.fixture(autouse=True)
def _isolate_job_side_effects(db_factory, tmp_path, monkeypatch):
    # 생성 작업 스레드가 실제 로컬 DB와 data/ 폴더를 건드리지 않게 테스트 DB·임시 폴더로 돌린다.
    monkeypatch.setattr(wireframe_routes, "SessionLocal", db_factory)
    monkeypatch.setattr(wireframe_routes, "_WIRE_OUTPUT_DIR", tmp_path)


@pytest.fixture(autouse=True)
def _reset_state():
    wireframe_routes._jobs.clear()
    wireframe_routes._themes.clear()


@pytest.fixture()
def admin(client, make_user, monkeypatch):
    monkeypatch.setattr(deps_module, "ADMIN_EMAIL", "admin@team.com")
    make_user(email="admin@team.com", password="correcthorse", name="관리자")
    client.post("/auth/login", json={"email": "admin@team.com", "password": "correcthorse"})


def _post(client, path, content, **data):
    return client.post(path, files={"file": ("sb.pptx", content)}, data=data)


def test_analyze_reports_whether_a_settings_slide_was_found(client, admin):
    with_settings = _deck(lambda p: _add_text_slide(p, GUIDE_TEXT), _add_screen_slide)
    body = _post(client, "/wireframe/analyze", with_settings).json()
    assert body["settings"]["slide"] == 1
    assert "레이아웃: 좌측 사이드바형" in body["settings"]["text"]

    without = _post(client, "/wireframe/analyze", _deck(_add_screen_slide)).json()
    assert without["settings"] is None


def test_settings_slide_reaches_theme_and_page_generation(client, admin, monkeypatch):
    seen = {}

    def fake_theme(context, on_text=None):
        seen["theme"] = context
        return ":root{}", {"input": 1, "output": 1}

    def fake_page(page, context, theme_css, related, instruction="", on_text=None):
        seen["page"] = context
        return "<!doctype html><html></html>", {"input": 1, "output": 1, "cacheRead": 0, "cacheWrite": 0}

    monkeypatch.setattr(wireframe_routes, "generate_theme", fake_theme)
    monkeypatch.setattr(wireframe_routes, "generate_page_hifi", fake_page)
    deck = _deck(lambda p: _add_text_slide(p, GUIDE_TEXT), _add_screen_slide)
    job_id = _post(client, "/wireframe/generate", deck, page_id="DN-SEA-P-0").json()["jobId"]
    for _ in range(400):
        if client.get(f"/wireframe/jobs/{job_id}").json()["status"] in ("done", "error"):
            break
        time.sleep(0.05)

    for key in ("theme", "page"):
        assert seen[key].startswith("[와이어프레임 설정 장표 — 최우선 반영]"), key
        assert "스타일: 모던 SaaS형" in seen[key], key


def test_screen_range_in_settings_slide_is_used_when_slides_have_no_page_ids(client, admin):
    deck = _deck(lambda p: _add_text_slide(p, GUIDE_TEXT), lambda p: _add_text_slide(p, ["화면 A"]), lambda p: _add_text_slide(p, ["화면 B"]))
    pages = _post(client, "/wireframe/analyze", deck).json()["pages"]
    assert [(page["pageId"], page["slides"]) for page in pages] == [("SLIDE-RANGE-2-3", [2, 3])]
