import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Pt

from app.wireframe.converter import convert
from app.wireframe.renderer import render_html


def _deck_bytes() -> io.BytesIO:
    prs = Presentation()
    prs.slide_width = Emu(12192000)
    prs.slide_height = Emu(6858000)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    button = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(952500), Emu(952500), Emu(1905000), Emu(476250))
    button.text_frame.text = "DOWNLOAD"
    button.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
    badge = slide.shapes.add_shape(MSO_SHAPE.OVAL, Emu(762000), Emu(762000), Emu(190500), Emu(190500))
    badge.text_frame.text = "#1"
    panel = slide.shapes.add_table(3, 2, Emu(9163900), Emu(472800), Emu(3024000), Emu(900000))
    panel.table.cell(0, 0).text = "#1"
    panel.table.cell(0, 1).text = "다운로드 버튼 - 클릭 시 런처 설치 파일 다운로드"
    panel.table.cell(1, 0).text = "#2"
    out = io.BytesIO()
    prs.save(out)
    out.seek(0)
    return out


def test_convert_places_shapes_badges_and_panel_notes():
    docs = convert(_deck_bytes())
    assert len(docs) == 1
    doc = docs[0]
    assert "DOWNLOAD" in doc.body
    assert 'data-n="1"' in doc.body
    assert [(a.number, a.text) for a in doc.annotations] == [(1, "다운로드 버튼 - 클릭 시 런처 설치 파일 다운로드")]
    assert "#2" not in doc.body


def test_render_html_links_badge_and_note_by_number():
    html_text = render_html(convert(_deck_bytes()), "테스트")
    assert html_text.count('data-n="1"') == 2
    assert "<title>테스트</title>" in html_text


def test_items_turn_soft_line_breaks_into_spaces():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Emu(0), Emu(0), Emu(3000000), Emu(500000))
    box.text_frame.paragraphs[0].text = "WHAT'S ON\vTHIS MONTH"
    out = io.BytesIO()
    prs.save(out)
    out.seek(0)
    texts = [item[5] for item in convert(out)[0].items]
    assert "WHAT'S ON THIS MONTH" in texts
